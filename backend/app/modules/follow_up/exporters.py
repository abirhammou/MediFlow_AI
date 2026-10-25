"""
Génération des exports : PDF, Word (.docx), Email (HTML professionnel).
"""
from __future__ import annotations

import io
from datetime import date


# ── PDF ───────────────────────────────────────────────────────────────────────

def generate_pdf(subject: str, body: str, patient_name: str,
                 appointment_date: str, appointment_reason: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    from reportlab.lib.enums import TA_CENTER, TA_LEFT

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=2.5 * cm, rightMargin=2.5 * cm,
        topMargin=2 * cm, bottomMargin=2 * cm,
    )

    styles = getSampleStyleSheet()
    PRIMARY = colors.HexColor("#1459d9")

    title_style = ParagraphStyle(
        "Title", parent=styles["Normal"],
        fontSize=18, fontName="Helvetica-Bold",
        textColor=PRIMARY, alignment=TA_LEFT, spaceAfter=6,
    )
    meta_style = ParagraphStyle(
        "Meta", parent=styles["Normal"],
        fontSize=9, textColor=colors.HexColor("#64748b"), spaceAfter=4,
    )
    subject_style = ParagraphStyle(
        "Subject", parent=styles["Normal"],
        fontSize=13, fontName="Helvetica-Bold", spaceAfter=12,
    )
    body_style = ParagraphStyle(
        "Body", parent=styles["Normal"],
        fontSize=11, leading=18, spaceAfter=8,
    )

    story = [
        Paragraph("MediFlow AI", title_style),
        Paragraph("Rappel administratif généré automatiquement", meta_style),
        HRFlowable(width="100%", thickness=1, color=PRIMARY, spaceAfter=16),
        Paragraph(f"Patient : <b>{patient_name}</b>", meta_style),
        Paragraph(f"Date du rendez-vous : <b>{appointment_date}</b>", meta_style),
        Paragraph(f"Motif : <b>{appointment_reason}</b>", meta_style),
        Spacer(1, 0.5 * cm),
        HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#e2e8f0"), spaceAfter=12),
        Paragraph(f"Objet : {subject}", subject_style),
    ]

    for line in body.split("\n"):
        if line.strip():
            story.append(Paragraph(line, body_style))
        else:
            story.append(Spacer(1, 0.3 * cm))

    story += [
        Spacer(1, 1 * cm),
        HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#e2e8f0"), spaceAfter=8),
        Paragraph(
            f"Document généré le {date.today().strftime('%d/%m/%Y')} · MediFlow AI",
            meta_style,
        ),
    ]

    doc.build(story)
    return buffer.getvalue()


# ── WORD ─────────────────────────────────────────────────────────────────────

def generate_docx(subject: str, body: str, patient_name: str,
                  appointment_date: str, appointment_reason: str) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    document = Document()

    # Marges
    for section in document.sections:
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)

    PRIMARY = RGBColor(0x14, 0x59, 0xD9)
    MUTED = RGBColor(0x64, 0x74, 0x8B)

    # En-tête
    h = document.add_heading("MediFlow AI", level=1)
    h.runs[0].font.color.rgb = PRIMARY

    sub = document.add_paragraph("Rappel administratif généré automatiquement")
    sub.runs[0].font.color.rgb = MUTED
    sub.runs[0].font.size = Pt(9)

    document.add_paragraph("─" * 60)

    # Métadonnées
    for label, value in [
        ("Patient", patient_name),
        ("Date du rendez-vous", appointment_date),
        ("Motif", appointment_reason),
    ]:
        p = document.add_paragraph()
        run_lbl = p.add_run(f"{label} : ")
        run_lbl.font.color.rgb = MUTED
        run_lbl.font.size = Pt(9)
        run_val = p.add_run(value)
        run_val.bold = True
        run_val.font.size = Pt(9)

    document.add_paragraph()

    # Objet
    obj_p = document.add_paragraph()
    obj_p.add_run("Objet : ").bold = True
    obj_p.add_run(subject)

    document.add_paragraph()

    # Corps
    for line in body.split("\n"):
        p = document.add_paragraph(line if line.strip() else " ")
        p.paragraph_format.space_after = Pt(4)

    document.add_paragraph()
    document.add_paragraph("─" * 60)

    footer_p = document.add_paragraph(
        f"Document généré le {date.today().strftime('%d/%m/%Y')} · MediFlow AI"
    )
    footer_p.runs[0].font.color.rgb = MUTED
    footer_p.runs[0].font.size = Pt(8)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ── EMAIL HTML ────────────────────────────────────────────────────────────────

def generate_email_html(subject: str, body: str, patient_name: str,
                        appointment_date: str, appointment_reason: str) -> str:
    body_html = "".join(
        f"<p style='margin:0 0 8px'>{line}</p>" if line.strip() else "<br/>"
        for line in body.split("\n")
    )
    return f"""<!DOCTYPE html>
<html lang="fr">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width,initial-scale=1">
  <title>{subject}</title>
</head>
<body style="margin:0;padding:0;background:#f4f7fc;font-family:Inter,Arial,sans-serif">
  <table width="100%" cellpadding="0" cellspacing="0">
    <tr><td align="center" style="padding:32px 16px">
      <table width="600" cellpadding="0" cellspacing="0"
             style="background:#fff;border-radius:16px;box-shadow:0 4px 24px rgba(20,89,217,.1);overflow:hidden;max-width:600px">

        <!-- Header -->
        <tr>
          <td style="background:#1459d9;padding:24px 32px">
            <p style="margin:0;color:#fff;font-size:20px;font-weight:800">MediFlow AI</p>
            <p style="margin:4px 0 0;color:rgba(255,255,255,.7);font-size:12px">
              Rappel administratif
            </p>
          </td>
        </tr>

        <!-- Meta -->
        <tr>
          <td style="padding:20px 32px 0;background:#f8fafc;border-bottom:1px solid #e2e8f0">
            <table width="100%">
              <tr>
                <td style="font-size:12px;color:#64748b;padding-bottom:6px">
                  <b>Patient :</b> {patient_name}
                </td>
                <td style="font-size:12px;color:#64748b;padding-bottom:6px;text-align:right">
                  <b>RDV :</b> {appointment_date}
                </td>
              </tr>
              <tr>
                <td colspan="2" style="font-size:12px;color:#64748b;padding-bottom:16px">
                  <b>Motif :</b> {appointment_reason}
                </td>
              </tr>
            </table>
          </td>
        </tr>

        <!-- Subject -->
        <tr>
          <td style="padding:24px 32px 8px">
            <p style="margin:0;font-size:16px;font-weight:700;color:#0f1f3d">
              {subject}
            </p>
          </td>
        </tr>

        <!-- Body -->
        <tr>
          <td style="padding:8px 32px 32px;font-size:14px;color:#334155;line-height:1.7">
            {body_html}
          </td>
        </tr>

        <!-- Footer -->
        <tr>
          <td style="padding:16px 32px;background:#f8fafc;border-top:1px solid #e2e8f0">
            <p style="margin:0;font-size:11px;color:#94a3b8;text-align:center">
              Document généré le {date.today().strftime("%d/%m/%Y")} · MediFlow AI
            </p>
          </td>
        </tr>

      </table>
    </td></tr>
  </table>
</body>
</html>"""
